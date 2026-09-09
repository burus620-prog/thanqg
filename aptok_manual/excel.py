"""
매뉴얼 이미지 관리표 만들기

원본 yaml 과 촬영 경로를 읽어 엑셀 표를 만든다.
이미 만들어 둔 표가 있으면 손으로 채운 칸(우선순위·담당·메모·이미지 있음)을
이미지 파일명을 기준으로 그대로 옮겨 온다. 다시 만들어도 적어 둔 것이 사라지지 않는다.

쓰는 법
    python excel.py                     이 폴더의 서비스 하나만
    python excel.py ../다른폴더 .        여러 폴더를 한 파일에
    python excel.py --out 이름.xlsx     파일 이름 지정

만들어지는 것
    매뉴얼-이미지관리표.xlsx
      · 이미지 목록 — 한 줄이 이미지 한 장
      · 요약        — 서비스·대분류별 진행률 (자동 집계)
      · 사용법      — 채우는 방법과 촬영 명령
"""

import sys, os, glob, argparse
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="2F5597")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GRAY = PatternFill("solid", fgColor="F2F2F2")

COLS = [
    ("번호", 6), ("서비스", 10), ("대분류", 14), ("화면명", 20), ("메뉴 경로", 34),
    ("화면ID", 22), ("흐름", 12), ("단계", 6), ("무엇을 담는 화면인가", 42),
    ("이미지 파일명", 30), ("저장 위치", 38), ("촬영 방식", 10),
    ("이미지 있음", 10), ("상태", 12), ("가림 필요", 9), ("설명 상태", 11),
    ("우선순위", 9), ("담당", 10), ("교체 사유 · 메모", 34),
]
KEEP_COLS = {"이미지 있음": 13, "우선순위": 17, "담당": 18, "교체 사유 · 메모": 19}
HEADER_ROW = 4


def service_name(base):
    """site.yaml 의 제목에서 서비스 이름을 뽑는다."""
    p = os.path.join(base, "site.yaml")
    if os.path.exists(p):
        cfg = yaml.safe_load(open(p, encoding="utf-8")) or {}
        title = str(cfg.get("title", ""))
        for tail in (" 사용자 매뉴얼", " 매뉴얼"):
            if title.endswith(tail):
                return title[: -len(tail)]
        if title:
            return title
    return os.path.basename(os.path.abspath(base))


def capture_method(base):
    """이미지마다 주소 이동인지 클릭인지 알아낸다."""
    how = {}
    p = os.path.join(base, "routes.yaml")
    if not os.path.exists(p):
        return how
    raw = open(p, encoding="utf-8").read()
    for name in set(__import__("re").findall(r"\$\{(\w+)\}", raw)):
        raw = raw.replace("${%s}" % name, "X")
    try:
        cfg = yaml.safe_load(raw) or {}
    except Exception:
        return how
    for r in cfg.get("routes") or []:
        steps = r.get("steps") or []
        mode = "주소" if any("url" in s for s in steps) else "클릭"
        for s in steps:
            if "shot" in s:
                how[s["shot"]] = mode
    return how


def collect(base):
    """원본 yaml 에서 이미지가 필요한 자리를 모두 찾는다."""
    rows = []
    how = capture_method(base)
    service = service_name(base)
    src = os.path.join(base, "src")
    for f in sorted(glob.glob(os.path.join(src, "*.yaml"))):
        for doc in yaml.safe_load_all(open(f, encoding="utf-8")):
            for item in doc or []:
                sid = item.get("screen_id")
                if not sid:
                    continue
                path = str(item.get("path", ""))
                top = path.split(">")[0].strip()
                needs_mask = bool(item.get("privacy_note"))
                flows = [("본 흐름", item.get("steps") or [])]
                for sf in item.get("sub_flows") or []:
                    flows.append((sf.get("name", "하위 흐름"), sf.get("steps") or []))
                for flow_name, steps in flows:
                    for st in steps:
                        img = st.get("image")
                        if not img:
                            continue
                        rows.append({
                            "서비스": service,
                            "대분류": top,
                            "화면명": item.get("screen_name"),
                            "메뉴 경로": path,
                            "화면ID": sid,
                            "흐름": flow_name,
                            "단계": st.get("no"),
                            "무엇을 담는 화면인가": (st.get("action") or "").strip(),
                            "이미지 파일명": img,
                            "저장 위치": "captures/%s/%s" % (sid, img),
                            "촬영 방식": how.get(img, "미지정"),
                            "가림 필요": "예" if needs_mask else "",
                            "설명 상태": "뼈대" if item.get("draft") else "작성완료",
                            "_base": base,
                        })
    return rows


def check_files(rows):
    """captures 폴더를 뒤져 이미지가 실제로 있는지 확인한다."""
    found = set()
    for base in {r["_base"] for r in rows}:
        cap = os.path.join(base, "captures")
        if not os.path.isdir(cap):
            continue
        for root, _dirs, files in os.walk(cap):
            for name in files:
                if name.lower().endswith(".png"):
                    found.add(name)
    return found


def load_previous(path):
    """전에 만든 표에서 손으로 채운 값을 읽어 온다."""
    kept = {}
    if not path or not os.path.exists(path):
        return kept
    try:
        wb = load_workbook(path, data_only=True)
        if "이미지 목록" not in wb.sheetnames:
            return kept
        ws = wb["이미지 목록"]
        for row in ws.iter_rows(min_row=HEADER_ROW + 1):
            key = row[9].value          # 이미지 파일명
            if not key:
                continue
            kept[key] = {name: row[idx - 1].value for name, idx in KEEP_COLS.items()}
    except Exception as e:
        print("  ! 이전 표를 읽지 못했습니다 (%s). 빈 값으로 시작합니다." % type(e).__name__)
    return kept


def build(rows, out_path, kept, found):
    wb = Workbook()
    nf = Font(name=FONT, size=10)
    hf = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    bold = Font(name=FONT, bold=True, size=10)
    thin = Side(style="thin", color="BFBFBF")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(vertical="center", wrap_text=True)
    cen = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # ── 이미지 목록
    ws = wb.active
    ws.title = "이미지 목록"
    ws["A1"] = "매뉴얼 이미지 관리표"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = ("노란 칸만 직접 채우면 됩니다. [이미지 있음] 에 O 를 넣으면 상태가 자동으로 바뀝니다. "
                "회색 칸은 원본 yaml 에서 가져온 값이라 고치지 마세요.")
    ws["A2"].font = Font(name=FONT, size=9, color="808080")
    ws.merge_cells("A2:S2")

    for i, (h, w) in enumerate(COLS, 1):
        c = ws.cell(HEADER_ROW, i, h)
        c.fill, c.font, c.border, c.alignment = HEAD_FILL, hf, bd, cen
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[HEADER_ROW].height = 30

    for n, rw in enumerate(rows, start=1):
        r = HEADER_ROW + n
        prev = kept.get(rw["이미지 파일명"], {})
        has = prev.get("이미지 있음") or ("O" if rw["이미지 파일명"] in found else "")
        vals = [n, rw["서비스"], rw["대분류"], rw["화면명"], rw["메뉴 경로"], rw["화면ID"],
                rw["흐름"], rw["단계"], rw["무엇을 담는 화면인가"], rw["이미지 파일명"],
                rw["저장 위치"], rw["촬영 방식"], has, None, rw["가림 필요"],
                rw["설명 상태"], prev.get("우선순위"), prev.get("담당"),
                prev.get("교체 사유 · 메모")]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.font, c.border = nf, bd
            if i in (5, 9, 19):
                c.alignment = wrap
            elif i in (1, 7, 8, 12, 13, 14, 15, 16, 17, 18):
                c.alignment = cen
            else:
                c.alignment = Alignment(vertical="center")
            c.fill = YELLOW if i in (13, 17, 18, 19) else GRAY
        ws.cell(r, 14).value = '=IF(M%d="","없음","있음")' % r
        ws.cell(r, 14).fill = GRAY
        ws.row_dimensions[r].height = 28

    last = HEADER_ROW + len(rows)
    dv = DataValidation(type="list", formula1='"O"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("M%d:M%d" % (HEADER_ROW + 1, last))
    dv2 = DataValidation(type="list", formula1='"높음,보통,낮음"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add("Q%d:Q%d" % (HEADER_ROW + 1, last))

    rng = "N%d:N%d" % (HEADER_ROW + 1, last)
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=['$N%d="없음"' % (HEADER_ROW + 1)],
        fill=PatternFill("solid", fgColor="FFC7CE"),
        font=Font(name=FONT, size=10, color="9C0006")))
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=['$N%d="있음"' % (HEADER_ROW + 1)],
        fill=PatternFill("solid", fgColor="C6EFCE"),
        font=Font(name=FONT, size=10, color="006100")))
    ws.conditional_formatting.add("O%d:O%d" % (HEADER_ROW + 1, last), FormulaRule(
        formula=['$O%d="예"' % (HEADER_ROW + 1)],
        fill=PatternFill("solid", fgColor="FCE4D6"),
        font=Font(name=FONT, size=10, color="C55A11")))
    ws.freeze_panes = "A%d" % (HEADER_ROW + 1)
    ws.auto_filter.ref = "A%d:S%d" % (HEADER_ROW, last)

    # ── 요약
    s = wb.create_sheet("요약")
    s["A1"] = "진행 현황"
    s["A1"].font = Font(name=FONT, bold=True, size=14)
    s["A2"] = "이 표는 [이미지 목록] 을 자동으로 집계합니다. 직접 고치지 마세요."
    s["A2"].font = Font(name=FONT, size=9, color="808080")
    for i, h in enumerate(["구분", "전체", "있음", "없음", "진행률"], 1):
        c = s.cell(4, i, h)
        c.fill, c.font, c.border, c.alignment = HEAD_FILL, hf, bd, cen
        s.column_dimensions[get_column_letter(i)].width = [22, 10, 10, 10, 12][i - 1]

    L = "'이미지 목록'"
    lo, hi = HEADER_ROW + 1, last
    pairs = []
    for rw in rows:
        k = (rw["서비스"], rw["대분류"])
        if k not in pairs:
            pairs.append(k)
    services = []
    for rw in rows:
        if rw["서비스"] not in services:
            services.append(rw["서비스"])

    row = 5
    s.cell(row, 1, "전체").font = bold
    s.cell(row, 2, "=COUNTA(%s!J%d:J%d)" % (L, lo, hi))
    s.cell(row, 3, '=COUNTIF(%s!N%d:N%d,"있음")' % (L, lo, hi))
    s.cell(row, 4, '=COUNTIF(%s!N%d:N%d,"없음")' % (L, lo, hi))
    s.cell(row, 5, "=IF(B%d=0,0,C%d/B%d)" % (row, row, row))
    row += 1
    for sv in services:
        s.cell(row, 1, sv).font = bold
        s.cell(row, 2, '=COUNTIF(%s!B%d:B%d,A%d)' % (L, lo, hi, row))
        s.cell(row, 3, '=COUNTIFS(%s!B%d:B%d,A%d,%s!N%d:N%d,"있음")' % (L, lo, hi, row, L, lo, hi))
        s.cell(row, 4, '=COUNTIFS(%s!B%d:B%d,A%d,%s!N%d:N%d,"없음")' % (L, lo, hi, row, L, lo, hi))
        s.cell(row, 5, "=IF(B%d=0,0,C%d/B%d)" % (row, row, row))
        row += 1
        for svc, top in pairs:
            if svc != sv:
                continue
            s.cell(row, 1, "   " + top)
            s.cell(row, 2, '=COUNTIFS(%s!B%d:B%d,"%s",%s!C%d:C%d,"%s")' % (L, lo, hi, sv, L, lo, hi, top))
            s.cell(row, 3, '=COUNTIFS(%s!B%d:B%d,"%s",%s!C%d:C%d,"%s",%s!N%d:N%d,"있음")'
                   % (L, lo, hi, sv, L, lo, hi, top, L, lo, hi))
            s.cell(row, 4, '=COUNTIFS(%s!B%d:B%d,"%s",%s!C%d:C%d,"%s",%s!N%d:N%d,"없음")'
                   % (L, lo, hi, sv, L, lo, hi, top, L, lo, hi))
            s.cell(row, 5, "=IF(B%d=0,0,C%d/B%d)" % (row, row, row))
            row += 1
    for rr in range(5, row):
        for cc in range(1, 6):
            c = s.cell(rr, cc)
            if not c.font.bold:
                c.font = nf
            c.border = bd
            if cc > 1:
                c.alignment = cen
        s.cell(rr, 5).number_format = "0.0%"

    # ── 사용법
    g = wb.create_sheet("사용법")
    g.column_dimensions["A"].width = 100
    guide = [
        ("매뉴얼 이미지 관리표 사용법", "T"), ("", ""),
        ("이 표는 무엇인가", "H"),
        ("매뉴얼에 들어갈 이미지를 한눈에 관리하는 표입니다.", "N"),
        ("어느 화면의 몇 번째 단계에 어떤 이미지가 필요한지, 그 이미지가 지금 있는지 없는지를 봅니다.", "N"),
        ("", ""),
        ("직접 채우는 칸 (노란색)", "H"),
        ("이미지 있음 — 파일을 넣었으면 O 를 고릅니다. 그러면 [상태] 가 '있음' 으로 바뀝니다.", "N"),
        ("우선순위 — 높음 / 보통 / 낮음. 어느 것부터 찍을지 정할 때 씁니다.", "N"),
        ("담당 — 누가 찍을지 적습니다.", "N"),
        ("교체 사유 · 메모 — 왜 바꿔야 하는지, 어떻게 찍어야 하는지 적습니다.", "N"),
        ("", ""),
        ("건드리지 않는 칸 (회색)", "H"),
        ("원본 yaml 에서 자동으로 가져온 값입니다. 여기를 고쳐도 매뉴얼은 바뀌지 않습니다.", "N"),
        ("매뉴얼 내용을 바꾸려면 src 폴더의 yaml 을 고치고 python build.py 를 실행하세요.", "N"),
        ("", ""),
        ("이미지를 직접 넣는 방법", "H"),
        ("1. [저장 위치] 에 적힌 경로에 파일을 둡니다.  예) captures/cs-all/cs-all-01.png", "N"),
        ("2. 파일 이름은 [이미지 파일명] 과 정확히 같아야 합니다.", "N"),
        ("3. python build.py 를 실행하면 매뉴얼에 붙습니다.", "N"),
        ("4. 이 표의 [이미지 있음] 에 O 를 넣어 둡니다.", "N"),
        ("", ""),
        ("표를 다시 만들 때", "H"),
        ("python excel.py 를 실행하면 최신 상태로 다시 만들어집니다.", "N"),
        ("이때 노란 칸에 적어 둔 값(우선순위·담당·메모)은 그대로 옮겨옵니다.", "N"),
        ("captures 폴더에 파일이 있으면 [이미지 있음] 이 자동으로 O 가 됩니다.", "N"),
        ("", ""),
        ("촬영 방식", "H"),
        ("주소 — capture.py 가 주소로 바로 이동해 찍습니다. 실패가 거의 없습니다.", "N"),
        ("클릭 — 메뉴를 눌러 이동합니다. 화면 구조가 바뀌면 실패할 수 있습니다.", "N"),
        ("", ""),
        ("가림 필요", "H"),
        ("'예' 인 화면은 개인정보나 금액이 나옵니다.", "N"),
        ("자동 촬영은 routes.yaml 의 mask_columns 설정으로 가려집니다.", "N"),
        ("직접 캡처해 넣을 때는 손으로 가린 뒤 저장하세요.", "N"),
        ("", ""),
        ("설명 상태", "H"),
        ("작성완료 — 화면 설명이 채워진 상태입니다.", "N"),
        ("뼈대 — 아직 짐작으로만 적혀 있어 화면을 보고 채워야 합니다.", "N"),
        ("", ""),
        ("자동 촬영 명령", "H"),
        ("python capture.py --headed              전체", "N"),
        ("python capture.py --headed --only cs    화면ID 에 cs 가 든 것만", "N"),
        ("python verify.py                        촬영본 검증 + review.html 생성", "N"),
        ("python build.py                         매뉴얼 생성", "N"),
    ]
    rr = 1
    for text, kind in guide:
        c = g.cell(rr, 1, text)
        if kind == "T":
            c.font = Font(name=FONT, bold=True, size=14)
        elif kind == "H":
            c.font = Font(name=FONT, bold=True, size=11, color="2F5597")
        else:
            c.font = nf
        c.alignment = Alignment(vertical="center", wrap_text=True)
        rr += 1

    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser(description="매뉴얼 이미지 관리표를 만든다")
    ap.add_argument("folders", nargs="*", default=["."],
                    help="서비스 폴더. 여러 개를 주면 한 파일에 모은다 (기본: 지금 폴더)")
    ap.add_argument("--out", default="매뉴얼-이미지관리표.xlsx", help="만들 파일 이름")
    args = ap.parse_args()

    folders = args.folders or ["."]
    rows = []
    for f in folders:
        if not os.path.isdir(os.path.join(f, "src")):
            print("  ! %s 에 src 폴더가 없습니다. 건너뜁니다." % f)
            continue
        got = collect(f)
        print("  %s : 이미지 %d장" % (service_name(f), len(got)))
        rows += got

    if not rows:
        print("\n이미지가 필요한 자리를 찾지 못했습니다.")
        print("서비스 폴더 안에서 실행하거나, 폴더 경로를 인자로 주세요.")
        return

    kept = load_previous(args.out)
    if kept:
        print("  이전 표에서 적어 둔 값 %d개를 옮겨 옵니다." % len(kept))
    found = check_files(rows)
    if found:
        print("  captures 폴더에서 이미지 %d장을 찾았습니다." % len(found))

    build(rows, args.out, kept, found)

    have = sum(1 for r in rows
               if kept.get(r["이미지 파일명"], {}).get("이미지 있음")
               or r["이미지 파일명"] in found)
    print("\n%s 를 만들었습니다." % args.out)
    print("  전체 %d장 · 있음 %d장 · 없음 %d장" % (len(rows), have, len(rows) - have))
    print("\n엑셀에서 열면 [요약] 탭의 진행률이 계산됩니다.")


if __name__ == "__main__":
    main()
